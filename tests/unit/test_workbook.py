import io

import pytest
from openpyxl import load_workbook

from phi_framework import constants as c
from phi_framework.workbook import DATA_SHEET, build_workbook, content_hash, parse_workbook

ROWS = [
    {"review_item_id": "ITEM-1", "catalog_name": "cat", "schema_name": "sch", "table_name": "t", "column_name": "a",
     "issue_type": c.MISSING_CLASSIFICATION, "status": c.PENDING_STEWARD_REVIEW, "assigned_steward": "s@example.com",
     "reviewed_by": "s@example.com"},
    {"review_item_id": "ITEM-2", "catalog_name": "cat", "schema_name": "sch", "table_name": "t", "column_name": "b",
     "issue_type": c.MISSING_CLASSIFICATION, "status": c.PENDING_STEWARD_REVIEW, "assigned_steward": "s@example.com",
     "reviewed_by": "s@example.com"},
]


def _build(wtype=c.WB_REVIEW):
    return build_workbook(wtype, "BATCH-1", "s@example.com", ROWS, ["mrn", "name"], {"env": "dev"}, "pw")


def test_round_trip_reads_only_editable_fields():
    data = _build()
    wb = load_workbook(io.BytesIO(data))
    ws = wb[DATA_SHEET]
    headers = [c_.value for c_ in ws[1]]
    ws.cell(row=2, column=headers.index("decision") + 1, value="APPROVE_SUGGESTION")
    ws.cell(row=2, column=headers.index("reviewed_at") + 1, value="2026-10-14 11:40 -04:00")
    ws.cell(row=2, column=headers.index("column_name") + 1, value="tampered")  # locked column: ignored
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue())
    assert parsed.workbook_type == c.WB_REVIEW and parsed.review_batch_id == "BATCH-1"
    assert [r.review_item_id for r in parsed.rows] == ["ITEM-1", "ITEM-2"]
    assert parsed.rows[0].values["decision"] == "APPROVE_SUGGESTION"
    assert parsed.rows[0].values["reviewed_at"] == "2026-10-14 11:40 -04:00"
    assert "column_name" not in parsed.rows[0].values
    assert parsed.rows[1].values["decision"] is None


def test_protection_and_text_format():
    ws = load_workbook(io.BytesIO(_build()))[DATA_SHEET]
    headers = [c_.value for c_ in ws[1]]
    assert ws.protection.sheet
    assert ws.cell(row=2, column=headers.index("review_item_id") + 1).protection.locked
    cell = ws.cell(row=2, column=headers.index("reviewed_at") + 1)
    assert not cell.protection.locked and cell.number_format == "@"


def test_assignment_layout_round_trip():
    parsed = parse_workbook(_build(c.WB_ASSIGNMENT))
    assert parsed.workbook_type == c.WB_ASSIGNMENT
    assert set(parsed.rows[0].values) == {"assigned_steward"}


def test_rejects_foreign_workbook():
    from openpyxl import Workbook

    buf = io.BytesIO()
    Workbook().save(buf)
    with pytest.raises(ValueError):
        parse_workbook(buf.getvalue())


def test_content_hash_changes_with_rows():
    assert content_hash("REVIEW", "B", "a", ROWS) != content_hash("REVIEW", "B", "a", ROWS[:1])
