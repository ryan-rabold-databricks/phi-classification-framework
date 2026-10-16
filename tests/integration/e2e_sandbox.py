"""End-to-end test of the deployed framework against synthetic data in a dev sandbox.

Drives the real jobs (setup, scan, ingestion, application) with run-now,
edits the exported workbooks the way Data Governance and stewards would, and
asserts every state transition. Synthetic data only.

  uv run --with openpyxl python tests/integration/e2e_sandbox.py \
      -p <profile> --warehouse-id <id> --test-schema-prefix phifw_test --reset

Bundle variables (governance catalog/schema, tag key, groups, ...) are read
from `databricks bundle validate -t dev`. The current user must be a member of
the steward group and the governance group.

--reset drops and recreates the governance schema and the synthetic schemas.
It refuses to run unless the target's env is dev.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import subprocess
import sys
import time
from datetime import datetime, timedelta, timezone

ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
sys.path.insert(0, os.path.join(ROOT, "src"))
sys.path.insert(0, os.path.join(ROOT, "tools"))

from databricks.sdk import WorkspaceClient
from databricks.sdk.service.jobs import RunResultState
from openpyxl import load_workbook
from sqlexec import Sql

from phi_framework import constants as c
from phi_framework.keys import make_key
from phi_framework.sqlutil import fqn, lit, qi
from phi_framework.workbook import DATA_SHEET

RESULTS: list[tuple[str, bool, str]] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(cond), detail))
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}{(' -- ' + detail) if detail and not cond else ''}")
    if not cond:
        raise AssertionError(f"{name}: {detail}")


class E2E:
    def __init__(self, a):
        self.a = a
        self.w = WorkspaceClient(profile=a.profile)
        self.sql = Sql(self.w, a.warehouse_id)
        self.me = self.w.current_user.me().user_name.lower()
        out = subprocess.run(
            ["databricks", "bundle", "validate", "-t", "dev", "-p", a.profile, "-o", "json"],
            check=True, capture_output=True, text=True, cwd=ROOT,
        ).stdout
        b = json.loads(out)
        self.v = {k: v.get("value") for k, v in b["variables"].items()}
        if self.v["env"] != "dev":
            raise SystemExit("E2E runs only against env=dev")
        summary = json.loads(
            subprocess.run(["databricks", "bundle", "summary", "-t", "dev", "-p", a.profile, "-o", "json"],
                           check=True, capture_output=True, text=True, cwd=ROOT).stdout
        )
        self.jobs = {k: int(j["id"]) for k, j in summary["resources"]["jobs"].items()}
        self.gc, self.gs = self.v["governance_catalog"], self.v["governance_schema"]
        self.cat = self.gc  # sandbox: synthetic schemas live in the governance catalog
        self.tag = self.v["tag_key"]
        self.clinical = f"{a.test_schema_prefix}_clinical"
        self.finance = f"{a.test_schema_prefix}_finance"
        self.support = f"{a.test_schema_prefix}_support"
        self.vol = f"/Volumes/{self.gc}/{self.gs}/{c.V_FILES}"

    # ---------------- helpers ----------------
    def t(self, name: str) -> str:
        return fqn(self.gc, self.gs, name)

    def q(self, s: str) -> list[dict]:
        return self.sql.run(s)

    def run_job(self, key: str, params: dict | None = None) -> int:
        waiter = self.w.jobs.run_now(job_id=self.jobs[key], job_parameters=params or {})
        run_id = waiter.response.run_id
        while True:
            run = self.w.jobs.get_run(run_id)
            if run.state.life_cycle_state.value in ("TERMINATED", "INTERNAL_ERROR", "SKIPPED"):
                break
            time.sleep(10)
        if run.state.result_state != RunResultState.SUCCESS:
            for t in run.tasks or []:
                if t.state.result_state != RunResultState.SUCCESS:
                    out = self.w.jobs.get_run_output(t.run_id)
                    print(f"  task {t.task_key} failed: {out.error}\n{(out.error_trace or '')[-3000:]}")
            raise RuntimeError(f"{key} run {run_id} finished {run.state.result_state}: {run.state.state_message}")
        print(f"  ran {key} (run {run_id})")
        return run_id

    def item(self, schema: str, table: str, column: str, open_only: bool = False) -> dict | None:
        cond = f"AND status IN ({', '.join(lit(s) for s in c.NON_TERMINAL_STATUSES)})" if open_only else ""
        rows = self.q(
            f"""SELECT * FROM {self.t(c.T_ITEM)} WHERE catalog_name = {lit(self.cat)} AND schema_name = {lit(schema)}
            AND table_name = {lit(table)} AND column_name = {lit(column)} {cond}
            ORDER BY audit_created_at DESC LIMIT 1"""
        )
        return rows[0] if rows else None

    def tag_value(self, schema: str, table: str, column: str) -> str | None:
        rows = self.q(
            f"""SELECT tag_value FROM {qi(self.cat)}.information_schema.column_tags WHERE schema_name = {lit(schema)}
            AND table_name = {lit(table)} AND column_name = {lit(column)} AND tag_name = {lit(self.tag)}"""
        )
        return rows[0]["tag_value"] if rows else None

    def events(self, event_type: str) -> list[dict]:
        return self.q(f"SELECT * FROM {self.t(c.T_EVENT)} WHERE event_type = {lit(event_type)}")

    def outbound(self, wtype: str) -> list[str]:
        paths = []
        for d in self.w.files.list_directory_contents(f"{self.vol}/outbound"):
            if d.is_directory:
                for f in self.w.files.list_directory_contents(d.path):
                    if f"__{wtype}__" in f.name:
                        paths.append((f.last_modified, f.path))
        return [p for _, p in sorted(paths)]

    def edit_and_upload(self, src_path: str, edits: dict, dest_name: str) -> str:
        """edits: {(schema, table, column) | review_item_id: {column: value}}"""
        data = self.w.files.download(src_path).contents.read()
        wb = load_workbook(io.BytesIO(data))
        ws = wb[DATA_SHEET]
        headers = [cell.value for cell in ws[1]]
        col = {h: i + 1 for i, h in enumerate(headers)}
        applied = 0
        for r in range(2, ws.max_row + 1):
            row = {h: ws.cell(row=r, column=col[h]).value for h in headers}
            keys = [row["review_item_id"], (row.get("schema_name"), row.get("table_name"), row.get("column_name"))]
            for k in keys:
                if k in edits:
                    for h, v in edits[k].items():
                        ws.cell(row=r, column=col[h], value=v)
                    applied += 1
        buf = io.BytesIO()
        wb.save(buf)
        dest = f"{self.vol}/inbound/{dest_name}"
        self.w.files.upload(dest, io.BytesIO(buf.getvalue()), overwrite=False)
        print(f"  uploaded {dest} ({applied} edited rows)")
        return dest

    @staticmethod
    def now_local(offset_minutes: int = -60) -> str:
        t = datetime.now(timezone(timedelta(hours=-4))) + timedelta(minutes=offset_minutes)
        return t.strftime("%Y-%m-%d %H:%M -04:00")

    # ---------------- setup ----------------
    def reset(self) -> None:
        print("== Reset sandbox")
        for sch in (self.gs, self.clinical, self.finance, self.support):
            self.q(f"DROP SCHEMA IF EXISTS {fqn(self.cat, sch)} CASCADE")
        # The jobs already exist, so the (faster) setup job can recreate the governance objects.
        self.run_job("phi_framework_setup")

    def seed(self) -> None:
        print("== Seed synthetic data")
        C = self.cat
        for sch in (self.clinical, self.finance, self.support):
            self.q(f"CREATE SCHEMA IF NOT EXISTS {fqn(C, sch)} COMMENT 'Synthetic data for PHI framework E2E tests'")
        self.q(f"""CREATE OR REPLACE TABLE {fqn(C, self.clinical, 'patient_encounter')} (
            patient_name STRING COMMENT 'Patient legal name', mrn STRING COMMENT 'Medical record number',
            encounter_date DATE COMMENT 'Date of encounter', source_system_code STRING COMMENT 'Technical source system identifier',
            notes STRING COMMENT 'Free text notes', ssn STRING COMMENT 'Social security number')""")
        self.q(f"""INSERT INTO {fqn(C, self.clinical, 'patient_encounter')} VALUES
            ('Jane Synthetic', 'MRN0001', DATE'2026-01-02', 'EPIC', 'none', '000-00-0000')""")
        self.q(f"""CREATE OR REPLACE TABLE {fqn(C, self.clinical, 'lab_result')} (
            patient_mrn STRING COMMENT 'MRN of the patient', result_value DOUBLE COMMENT 'Numeric result',
            legacy_flag STRING COMMENT 'Pre-tagged with a value that is not allowed')""")
        self.q(f"ALTER TABLE {fqn(C, self.clinical, 'lab_result')} ALTER COLUMN legacy_flag SET TAGS ({lit(self.tag)} = 'retired_value')")
        self.q(f"""CREATE OR REPLACE TABLE {fqn(C, self.finance, 'billing')} (
            account_number STRING COMMENT 'Billing account number', amount DECIMAL(10,2) COMMENT 'Charge amount',
            patient_email STRING COMMENT 'Patient email address')""")
        self.mv = False
        if not self.a.skip_mv:
            try:
                self.q(f"""CREATE OR REPLACE MATERIALIZED VIEW {fqn(C, self.clinical, 'encounter_summary_mv')}
                    COMMENT 'Synthetic MV' AS SELECT mrn, count(*) AS encounter_count
                    FROM {fqn(C, self.clinical, 'patient_encounter')} GROUP BY mrn""")
                self.mv = True
            except Exception as exc:
                print(f"  (materialized view skipped: {str(exc)[:200]})")
        # Synthetic Data Classification results (same schema as system.data_classification.results).
        dc = fqn(C, self.support, "dc_results")
        self.q(f"""CREATE OR REPLACE TABLE {dc} (latest_detected_time TIMESTAMP, first_detected_time TIMESTAMP,
            catalog_id STRING, table_id STRING, catalog_name STRING, schema_name STRING, table_name STRING,
            column_name STRING, data_type STRING, class_tag STRING, samples ARRAY<STRING>, confidence STRING, frequency FLOAT)""")
        rows = [
            (self.clinical, "patient_encounter", "patient_name", "class.name", "HIGH"),
            (self.clinical, "patient_encounter", "mrn", "class.medical_record_number", "HIGH"),
            (self.clinical, "patient_encounter", "ssn", "class.us_ssn", "HIGH"),
            (self.clinical, "patient_encounter", "notes", "class.name", "HIGH"),
            (self.clinical, "patient_encounter", "notes", "class.location", "HIGH"),  # conflict -> no suggestion
            (self.clinical, "patient_encounter", "encounter_date", "class.medical_date", "LOW"),  # low -> none
            (self.finance, "billing", "account_number", "class.bank_number", "HIGH"),
            (self.finance, "billing", "account_number", "class.credit_card", "HIGH"),  # same mapped value -> ok
            (self.finance, "billing", "patient_email", "class.email_address", "HIGH"),
        ]
        values = ", ".join(
            f"(current_timestamp(), current_timestamp(), 'x', 'y', {lit(C)}, {lit(s)}, {lit(t)}, {lit(col)}, 'string', "
            f"{lit(tag)}, array('SAMPLE-SECRET-VALUE'), {lit(conf)}, 0.9)"
            for s, t, col, tag, conf in rows
        )
        self.q(f"INSERT INTO {dc} VALUES {values}")
        stewards, gov = "phifw_clinical_stewards", self.v["governance_group"]
        scopes = f"""scope_id,catalog_name,schema_name,environment,risk_tier,steward_group,governance_owner,scan_enabled,effective_date,expiration_date
SCOPE-CLINICAL-DEV,{C},{self.clinical},dev,HIGH,{stewards},{gov},true,2026-01-01,
SCOPE-FINANCE-DEV,{C},{self.finance},dev,STANDARD,{stewards},{gov},true,2026-01-01,
"""
        path = "/tmp/phifw_e2e_scopes.csv"
        with open(path, "w") as fh:
            fh.write(scopes)
        subprocess.run(
            ["uv", "run", "--quiet", "python", "tools/load_reference_data.py", "-p", self.a.profile, "--warehouse-id",
             self.a.warehouse_id, "--catalog", self.gc, "--schema", self.gs, "--scopes", path,
             "--mapping", "reference_data/suggestion_mapping.csv"], check=True, cwd=ROOT,
        )
        self.stewards = stewards

    # ---------------- scenario ----------------
    def run(self) -> None:
        if self.a.reset:
            self.reset()
        self.seed()
        cl, fi = self.clinical, self.finance

        print("== 0. Key parity (Python vs SQL make_key)")
        sql_key = self.sql.scalar(f"SELECT {fqn(self.gc, self.gs, 'make_key')}('ITEM', array('a', NULL, 'c'))")
        check("make_key parity", sql_key == make_key("ITEM", ["a", None, "c"]), f"{sql_key}")

        print("== 1. Setup job is idempotent")
        self.run_job("phi_framework_setup")

        print("== 2. First scan")
        self.run_job("phi_classification_scan")
        batches = self.q(f"SELECT * FROM {self.t(c.T_BATCH)}")
        check("two batches (one per scope)", len(batches) == 2, str(len(batches)))
        pn = self.item(cl, "patient_encounter", "patient_name")
        check("MISSING_CLASSIFICATION opened", pn and pn["issue_type"] == c.MISSING_CLASSIFICATION)
        check("suggestion from DC", pn["suggested_tag_value"] == "name", str(pn["suggested_tag_value"]))
        check("conflicting detections -> no suggestion", self.item(cl, "patient_encounter", "notes")["suggested_tag_value"] is None)
        check("LOW confidence -> no suggestion", self.item(cl, "patient_encounter", "encounter_date")["suggested_tag_value"] is None)
        check("two HIGH detections mapping to one value -> suggestion",
              self.item(fi, "billing", "account_number")["suggested_tag_value"] == "financial")
        check("invalid pre-existing tag -> INVALID_TAG_VALUE",
              self.item(cl, "lab_result", "legacy_flag")["issue_type"] == c.INVALID_TAG_VALUE)
        leaked = self.sql.scalar(f"SELECT count(*) FROM {self.t(c.T_ITEM)} WHERE suggestion_reason LIKE '%SAMPLE-SECRET%'")
        check("DC sample values never copied", int(leaked) == 0)
        n_items = int(self.sql.scalar(f"SELECT count(*) FROM {self.t(c.T_ITEM)}"))
        check("every in-scope column has an item", n_items == (14 if self.mv else 12), str(n_items))
        assign = self.outbound(c.WB_ASSIGNMENT)
        check("assignment workbook per batch", len(assign) == 2, str(assign))

        print("== 3. Rerun scan with unchanged metadata opens nothing")
        self.run_job("phi_classification_scan")
        check("no new items or batches",
              int(self.sql.scalar(f"SELECT count(*) FROM {self.t(c.T_ITEM)}")) == n_items
              and int(self.sql.scalar(f"SELECT count(*) FROM {self.t(c.T_BATCH)}")) == 2)
        check("no duplicate assignment workbooks", len(self.outbound(c.WB_ASSIGNMENT)) == 2)

        print("== 4. Data Governance assigns stewards")
        bogus = "not.a.steward@example.com"
        for i, path in enumerate(assign):
            edits = {}
            data = self.w.files.download(path).contents.read()
            ws = load_workbook(io.BytesIO(data))[DATA_SHEET]
            hdr = [x.value for x in ws[1]]
            for r in range(2, ws.max_row + 1):
                iid = ws.cell(row=r, column=hdr.index("review_item_id") + 1).value
                colname = ws.cell(row=r, column=hdr.index("column_name") + 1).value
                edits[iid] = {"assigned_steward": bogus if colname == "amount" else self.me}
            self.edit_and_upload(path, edits, f"assign_{i}_v1.xlsx")
        self.run_job("phi_decision_ingestion")
        check("stewards assigned", self.item(cl, "patient_encounter", "patient_name")["assigned_steward"] == self.me)
        check("non-member assignment rejected", self.item(fi, "billing", "amount")["assigned_steward"] is None)
        check("STEWARD_ASSIGNED events", len(self.events(c.EV_STEWARD_ASSIGNED)) == n_items - 1)
        reviews = self.outbound(c.WB_REVIEW)
        check("review workbook per batch for the steward", len(reviews) == 2, str(reviews))
        check("re-export of assignment workbook for the rejected row", len(self.outbound(c.WB_ASSIGNMENT)) == 3)
        reassign = self.outbound(c.WB_ASSIGNMENT)[-1]
        self.edit_and_upload(reassign, {(fi, "billing", "amount"): {"assigned_steward": self.me}}, "assign_fin_v2.xlsx")

        print("== 5. Steward decisions")
        at = self.now_local()
        dec = {
            (cl, "patient_encounter", "patient_name"): {"decision": c.APPROVE_SUGGESTION, "reviewed_at": at},
            (cl, "patient_encounter", "mrn"): {"decision": c.APPROVE_SUGGESTION, "reviewed_at": at},
            (cl, "patient_encounter", "ssn"): {"decision": c.APPROVE_SUGGESTION, "reviewed_at": at},
            (cl, "patient_encounter", "encounter_date"): {"decision": c.CORRECT_CLASSIFICATION, "corrected_tag_value": "date", "reviewed_at": at},
            (cl, "patient_encounter", "source_system_code"): {"decision": c.CONFIRM_NOT_PHI, "steward_comment": "Technical source identifier only", "reviewed_at": at},
            (cl, "patient_encounter", "notes"): {"decision": c.REQUEST_PRIVACY_REVIEW, "steward_comment": "Free text may contain names", "reviewed_at": at},
            (cl, "lab_result", "patient_mrn"): {"decision": c.CORRECT_CLASSIFICATION, "corrected_tag_value": "not_a_value", "reviewed_at": at},
            (cl, "lab_result", "legacy_flag"): {"decision": c.CONFIRM_NOT_PHI, "steward_comment": "Legacy flag, not PHI", "reviewed_at": at},
            (cl, "encounter_summary_mv", "mrn"): {"decision": c.CORRECT_CLASSIFICATION, "corrected_tag_value": "mrn", "reviewed_at": at},
            (cl, "encounter_summary_mv", "encounter_count"): {"decision": c.CONFIRM_NOT_PHI, "steward_comment": "Aggregate count", "reviewed_at": at},
            (fi, "billing", "account_number"): {"decision": c.APPROVE_SUGGESTION, "reviewed_at": at},
            (fi, "billing", "patient_email"): {"decision": c.APPROVE_SUGGESTION, "reviewed_at": "10/14/2026 11:40"},
        }
        uploaded = {}
        for i, path in enumerate(reviews):
            uploaded[path] = self.edit_and_upload(path, dec, f"review_{i}_v1.xlsx")
        self.run_job("phi_decision_ingestion")
        check("approve -> PLANNED", self.item(cl, "patient_encounter", "patient_name")["status"] == c.PLANNED)
        check("correct -> PLANNED", self.item(cl, "patient_encounter", "encounter_date")["status"] == c.PLANNED)
        check("not PHI (untagged) -> ATTESTED_NOT_PHI", self.item(cl, "patient_encounter", "source_system_code")["status"] == c.ATTESTED_NOT_PHI)
        check("not PHI (tagged) -> PLANNED removal", self.item(cl, "lab_result", "legacy_flag")["status"] == c.PLANNED)
        check("privacy request -> NEEDS_PRIVACY_REVIEW", self.item(cl, "patient_encounter", "notes")["status"] == c.NEEDS_PRIVACY_REVIEW)
        check("invalid corrected value -> VALIDATION_FAILED", self.item(cl, "lab_result", "patient_mrn")["status"] == c.VALIDATION_FAILED)
        check("bad reviewed_at -> VALIDATION_FAILED", self.item(fi, "billing", "patient_email")["status"] == c.VALIDATION_FAILED)
        att = self.q(f"SELECT * FROM {self.t(c.T_ATTESTATION)} WHERE column_name = 'source_system_code'")[0]
        months = (datetime.fromisoformat(att["expires_at"][:19]) - datetime.fromisoformat(att["attested_at"][:19])).days
        check("HIGH-tier attestation expires in ~3 months", 88 <= months <= 93, str(months))
        plan = self.q(f"SELECT * FROM {self.t(c.T_PLAN)} WHERE column_name = 'legacy_flag'")[0]
        check("removal plan renders UNSET TAGS", plan["action_type"] == c.REMOVE_PHI_TAG and "UNSET TAGS" in plan["rendered_ddl"])
        if self.mv:
            mvp = self.q(f"SELECT rendered_ddl FROM {self.t(c.T_PLAN)} WHERE table_name = 'encounter_summary_mv'")[0]
            check("MV plan uses ALTER MATERIALIZED VIEW", mvp["rendered_ddl"].startswith("ALTER MATERIALIZED VIEW"))
        check("privacy workbook exported", len(self.outbound(c.WB_PRIVACY)) == 1)

        print("== 6. Plans are append-only")
        try:
            self.q(f"UPDATE {self.t(c.T_PLAN)} SET rendered_ddl = 'x'")
            check("UPDATE on plan rejected", False, "update succeeded")
        except RuntimeError as exc:
            check("UPDATE on plan rejected", "APPEND_ONLY" in str(exc).upper() or "append" in str(exc).lower(), str(exc)[:200])

        print("== 7. Tag application")
        self.run_job("phi_classification_application")
        check("tag applied and verified", self.item(cl, "patient_encounter", "patient_name")["status"] == c.TAG_VERIFIED
              and self.tag_value(cl, "patient_encounter", "patient_name") == "name")
        check("corrected value applied", self.tag_value(cl, "patient_encounter", "encounter_date") == "date")
        check("invalid tag removed -> ATTESTED_NOT_PHI", self.item(cl, "lab_result", "legacy_flag")["status"] == c.ATTESTED_NOT_PHI
              and self.tag_value(cl, "lab_result", "legacy_flag") is None)
        if self.mv:
            check("MV column tagged", self.tag_value(cl, "encounter_summary_mv", "mrn") == "mrn")

        print("== 8. Privacy outcome, rejected-row fixes, and an ignored change")
        priv = self.outbound(c.WB_PRIVACY)[-1]
        self.edit_and_upload(priv, {(cl, "patient_encounter", "notes"): {
            "decision": c.CONFIRM_NOT_PHI, "steward_comment": "Privacy: notes are system-generated", "reviewed_by": self.me,
            "reviewed_at": self.now_local(-5)}}, "privacy_v1.xlsx")
        # Revised copy of an already-accepted workbook changes an accepted decision -> ignored and reported.
        first_review = next(p for p in reviews if self.item(cl, "patient_encounter", "patient_name")["review_batch_id"] in p)
        self.edit_and_upload(first_review, {**dec, (cl, "patient_encounter", "patient_name"): {
            "decision": c.CORRECT_CLASSIFICATION, "corrected_tag_value": "mrn", "reviewed_at": at}}, "review_clinical_v2.xlsx")
        self.run_job("phi_decision_ingestion")
        check("privacy outcome -> ATTESTED_NOT_PHI", self.item(cl, "patient_encounter", "notes")["status"] == c.ATTESTED_NOT_PHI)
        check("change to accepted item ignored", len(self.events(c.EV_CHANGE_IGNORED)) >= 1
              and self.tag_value(cl, "patient_encounter", "patient_name") == "name")
        follow_ups = [p for p in self.outbound(c.WB_REVIEW) if p not in reviews]
        check("follow-up review workbooks exported", len(follow_ups) >= 1, str(follow_ups))
        fix = {
            (cl, "lab_result", "patient_mrn"): {"decision": c.CORRECT_CLASSIFICATION, "corrected_tag_value": "mrn", "reviewed_at": self.now_local(-2)},
            (cl, "lab_result", "result_value"): {"decision": c.CONFIRM_NOT_PHI, "steward_comment": "Lab value without identifiers", "reviewed_at": self.now_local(-2)},
            (fi, "billing", "patient_email"): {"decision": c.APPROVE_SUGGESTION, "reviewed_at": self.now_local(-2)},
            (fi, "billing", "amount"): {"decision": c.CONFIRM_NOT_PHI, "steward_comment": "Amount only", "reviewed_at": self.now_local(-2)},
        }
        for i, p in enumerate(self.outbound(c.WB_REVIEW)):
            if p not in reviews:
                self.edit_and_upload(p, fix, f"review_fix_{i}.xlsx")
        self.run_job("phi_decision_ingestion")
        self.run_job("phi_classification_application")
        open_items = int(self.sql.scalar(
            f"SELECT count(*) FROM {self.t(c.T_ITEM)} WHERE status NOT IN ({', '.join(lit(s) for s in c.TERMINAL_STATUSES)})"))
        check("all items terminal", open_items == 0, str(open_items))
        closed = int(self.sql.scalar(f"SELECT count(*) FROM {self.t(c.T_BATCH)} WHERE status = 'CLOSED'"))
        check("both batches CLOSED", closed == 2, str(closed))
        check("BATCH_CLOSED events", len(self.events(c.EV_BATCH_CLOSED)) == 2)

        print("== 9. Same workbook uploaded twice is processed once")
        before = len(self.events(c.EV_WORKBOOK_INGESTED))
        data = self.w.files.download(uploaded[reviews[0]]).contents.read()
        self.w.files.upload(f"{self.vol}/inbound/renamed_copy.xlsx", io.BytesIO(data), overwrite=False)
        self.run_job("phi_decision_ingestion")
        check("duplicate content skipped", len(self.events(c.EV_WORKBOOK_INGESTED)) == before)

        print("== 10. Metadata changes reopen items; removed columns close")
        self.q(f"ALTER TABLE {fqn(self.cat, cl, 'patient_encounter')} ALTER COLUMN source_system_code COMMENT 'Changed purpose'")
        self.q(f"ALTER TABLE {fqn(self.cat, cl, 'patient_encounter')} ALTER COLUMN ssn COMMENT 'SSN (changed)'")
        self.q(f"CREATE OR REPLACE TABLE {fqn(self.cat, fi, 'temp_drop')} (x STRING)")
        self.run_job("phi_classification_scan")
        check("STALE_NON_PHI_ATTESTATION", self.item(cl, "patient_encounter", "source_system_code", True)["issue_type"] == c.STALE_NON_PHI_ATTESTATION)
        check("STALE_PHI_REVIEW", self.item(cl, "patient_encounter", "ssn", True)["issue_type"] == c.STALE_PHI_REVIEW)
        check("new table -> MISSING_CLASSIFICATION", self.item(fi, "temp_drop", "x", True)["issue_type"] == c.MISSING_CLASSIFICATION)
        self.q(f"DROP TABLE {fqn(self.cat, fi, 'temp_drop')}")
        self.run_job("phi_classification_scan")
        check("removed column -> CLOSED_NO_ACTION", self.item(fi, "temp_drop", "x")["status"] == c.CLOSED_NO_ACTION)

        print("== 11. Correction request and tag removal on a verified PHI column")
        col = f"{self.cat}.{cl}.patient_encounter.mrn"
        self.run_job("phi_classification_scan", {"correction_column": col, "correction_reason": "E2E correction test",
                                                 "correction_requested_by": self.me})
        corr = self.item(cl, "patient_encounter", "mrn", True)
        check("CORRECTION_REQUESTED opened with current tag pre-filled",
              corr and corr["issue_type"] == c.CORRECTION_REQUESTED and corr["suggested_tag_value"] == "mrn")
        self.run_job("phi_classification_scan", {"correction_column": f"{self.cat}.{cl}.patient_encounter.ssn",
                                                 "correction_reason": "already open", "correction_requested_by": self.me})
        check("no second item when one is open", int(self.sql.scalar(
            f"SELECT count(*) FROM {self.t(c.T_ITEM)} WHERE column_name = 'ssn' AND status NOT IN ('TAG_VERIFIED','ATTESTED_NOT_PHI','CLOSED_NO_ACTION')")) == 1)
        # Assign and decide the correction: confirm not PHI on a tagged column -> removal.
        assign = [p for p in self.outbound(c.WB_ASSIGNMENT) if corr["review_batch_id"] in p][-1]
        self.edit_and_upload(assign, {corr["review_item_id"]: {"assigned_steward": self.me}}, "assign_corr.xlsx")
        self.run_job("phi_decision_ingestion")
        rev = [p for p in self.outbound(c.WB_REVIEW) if corr["review_batch_id"] in p][-1]
        self.edit_and_upload(rev, {corr["review_item_id"]: {"decision": c.CONFIRM_NOT_PHI,
                                   "steward_comment": "Correction: surrogate key only", "reviewed_at": self.now_local(-1)}}, "review_corr.xlsx")
        self.run_job("phi_decision_ingestion")
        self.run_job("phi_classification_application")
        check("correction removed tag and attested", self.item(cl, "patient_encounter", "mrn")["status"] == c.ATTESTED_NOT_PHI
              and self.tag_value(cl, "patient_encounter", "mrn") is None)
        prior = self.q(f"SELECT status FROM {self.t(c.T_ITEM)} WHERE column_name = 'mrn' AND table_name = 'patient_encounter' "
                       f"AND issue_type = 'MISSING_CLASSIFICATION'")
        check("original item unchanged as evidence", prior and prior[0]["status"] == c.TAG_VERIFIED)

        print("== 12. Audit columns")
        bad = self.sql.scalar(
            f"SELECT count(*) FROM {self.t(c.T_ITEM)} WHERE audit_created_process NOT LIKE 'dev_phi_%' OR audit_updated_run_id IS NULL")
        check("item audit columns populated by jobs", int(bad) == 0)
        try:
            self.q(f"UPDATE {self.t(c.T_ITEM)} SET audit_updated_process = 'MANUAL' WHERE true")
            check("manual process rejected on job-only table", False, "update succeeded")
        except RuntimeError as exc:
            check("manual process rejected on job-only table", "CHECK" in str(exc).upper(), str(exc)[:200])


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-p", "--profile", required=True)
    ap.add_argument("--warehouse-id", required=True)
    ap.add_argument("--test-schema-prefix", default="phifw_test")
    ap.add_argument("--reset", action="store_true")
    ap.add_argument("--skip-mv", action="store_true")
    a = ap.parse_args()
    try:
        E2E(a).run()
    finally:
        passed = sum(1 for _, ok, _ in RESULTS if ok)
        print(f"\n{passed}/{len(RESULTS)} checks passed")
    if any(not ok for _, ok, _ in RESULTS):
        sys.exit(1)


if __name__ == "__main__":
    main()
