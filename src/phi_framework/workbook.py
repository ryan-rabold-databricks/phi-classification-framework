"""Excel workbook layouts for the Phase 1 pilot.

Workbooks are built and parsed entirely in memory: Unity Catalog Volumes do
not support the random-access writes zip files need, so bytes are written to
the Volume in one call (spike S-05).

Sheet protection only prevents accidental edits; it is not a security
control. Ingestion never trusts locked columns -- identity and metadata always
come from the governance tables, keyed by review_item_id.
"""

from __future__ import annotations

import hashlib
import io
import json
from dataclasses import dataclass

from phi_framework import constants as c

META_SHEET = "_meta"
LISTS_SHEET = "_lists"
DATA_SHEET = "items"

# (header, width, editable)
REVIEW_COLUMNS = [
    ("review_item_id", 26, False),
    ("catalog_name", 24, False),
    ("schema_name", 18, False),
    ("table_name", 26, False),
    ("column_name", 24, False),
    ("full_data_type", 14, False),
    ("column_comment", 30, False),
    ("issue_type", 26, False),
    ("current_tag_value", 18, False),
    ("suggested_tag_value", 20, False),
    ("suggestion_reason", 30, False),
    ("status", 22, False),
    ("last_validation_error", 40, False),
    ("assigned_steward", 26, False),
    ("decision", 26, True),
    ("corrected_tag_value", 22, True),
    ("steward_comment", 40, True),
    ("reviewed_by", 28, True),
    ("reviewed_at", 22, True),
]

ASSIGNMENT_COLUMNS = [
    ("review_item_id", 26, False),
    ("scope_id", 22, False),
    ("steward_group", 24, False),
    ("catalog_name", 24, False),
    ("schema_name", 18, False),
    ("table_name", 26, False),
    ("column_name", 24, False),
    ("issue_type", 26, False),
    ("suggested_tag_value", 20, False),
    ("assigned_steward", 30, True),
]

EDITABLE_REVIEW = [h for h, _, e in REVIEW_COLUMNS if e]
MAX_ROWS = 20000


@dataclass(frozen=True)
class ParsedRow:
    row_number: int
    review_item_id: str
    values: dict  # editable column -> value (as text or None)


@dataclass(frozen=True)
class ParsedWorkbook:
    workbook_type: str
    review_batch_id: str
    meta: dict
    rows: list[ParsedRow]


def content_hash(workbook_type: str, review_batch_id: str, audience: str, rows: list[dict]) -> str:
    payload = json.dumps(
        {"t": workbook_type, "b": review_batch_id, "a": audience, "r": rows}, sort_keys=True, default=str
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def build_workbook(
    workbook_type: str,
    review_batch_id: str,
    audience: str,
    rows: list[dict],
    allowed_values: list[str],
    meta: dict,
    password: str,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation

    columns = ASSIGNMENT_COLUMNS if workbook_type == c.WB_ASSIGNMENT else REVIEW_COLUMNS
    if len(rows) > MAX_ROWS:
        raise ValueError(f"Workbook would contain {len(rows)} rows; the limit is {MAX_ROWS}. Split the batch.")
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    header_fill = PatternFill("solid", fgColor="1F3864")
    edit_fill = PatternFill("solid", fgColor="FFF2CC")
    for idx, (header, width, editable) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    for r, row in enumerate(rows, start=2):
        for idx, (header, _, editable) in enumerate(columns, start=1):
            value = row.get(header)
            cell = ws.cell(row=r, column=idx, value=None if value is None else str(value))
            if value is not None:
                cell.data_type = "s"  # store literally: metadata starting with '=' must never become a formula
            if editable:
                cell.protection = Protection(locked=False)
                cell.fill = edit_fill
                cell.number_format = "@"  # Text, so Excel never converts reviewed_at
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(columns)).column_letter}{max(len(rows) + 1, 2)}"

    lists = wb.create_sheet(LISTS_SHEET)
    for i, v in enumerate(sorted(allowed_values), start=1):
        lists.cell(row=i, column=1, value=v)
    for i, v in enumerate(c.DECISIONS, start=1):
        lists.cell(row=i, column=2, value=v)
    lists.sheet_state = "veryHidden"
    lists.protection.sheet = True
    lists.protection.password = password

    last = max(len(rows) + 1, 2)
    col_letter = {h: ws.cell(row=1, column=i).column_letter for i, (h, _, _) in enumerate(columns, start=1)}
    if workbook_type != c.WB_ASSIGNMENT:
        dv_decision = DataValidation(type="list", formula1=f"{LISTS_SHEET}!$B$1:$B${len(c.DECISIONS)}", allow_blank=True)
        dv_value = DataValidation(
            type="list", formula1=f"{LISTS_SHEET}!$A$1:$A${max(len(allowed_values), 1)}", allow_blank=True
        )
        for dv in (dv_decision, dv_value):
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
        dv_decision.add(f"{col_letter['decision']}2:{col_letter['decision']}{last}")
        dv_value.add(f"{col_letter['corrected_tag_value']}2:{col_letter['corrected_tag_value']}{last}")

    meta_ws = wb.create_sheet(META_SHEET)
    full_meta = {"workbook_type": workbook_type, "review_batch_id": review_batch_id, "audience": audience, **meta}
    for i, (k, v) in enumerate(full_meta.items(), start=1):
        meta_ws.cell(row=i, column=1, value=k)
        meta_ws.cell(row=i, column=2, value=str(v))
    meta_ws.sheet_state = "veryHidden"
    meta_ws.protection.sheet = True
    meta_ws.protection.password = password

    ws.protection.sheet = True
    ws.protection.password = password
    ws.protection.autoFilter = False
    ws.protection.sort = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _text(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def parse_workbook(data: bytes) -> ParsedWorkbook:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    if META_SHEET not in wb.sheetnames or DATA_SHEET not in wb.sheetnames:
        raise ValueError("Not a framework workbook: required sheets are missing")
    meta = {}
    for row in wb[META_SHEET].iter_rows(values_only=True):
        if row and row[0]:
            meta[str(row[0])] = None if row[1] is None else str(row[1])
    wtype = meta.get("workbook_type")
    batch = meta.get("review_batch_id")
    if wtype not in c.SUBMISSION_TYPES or not batch:
        raise ValueError(f"Workbook metadata is invalid: type={wtype!r}, batch={batch!r}")
    columns = ASSIGNMENT_COLUMNS if wtype == c.WB_ASSIGNMENT else REVIEW_COLUMNS
    expected = [h for h, _, _ in columns]
    ws = wb[DATA_SHEET]
    header = [(_text(c_.value) or "") for c_ in ws[1]][: len(expected)]
    if header != expected:
        raise ValueError(f"Workbook header does not match the {wtype} layout")
    idx = {h: i for i, h in enumerate(expected)}
    editable = [h for h, _, e in columns if e]
    rows: list[ParsedRow] = []
    seen: dict[str, int] = {}
    for r_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        item_id = _text(row[idx["review_item_id"]]) if row else None
        if not item_id:
            continue
        if item_id in seen:
            raise ValueError(f"review_item_id {item_id} appears on rows {seen[item_id]} and {r_num}; remove the duplicate")
        seen[item_id] = r_num
        values = {h: _text(row[idx[h]]) for h in editable}
        rows.append(ParsedRow(row_number=r_num, review_item_id=item_id, values=values))
    return ParsedWorkbook(workbook_type=wtype, review_batch_id=batch, meta=meta, rows=rows)
